"""Carga local y reproducible de fuentes oficiales CBAM a la biblioteca SICR3P.

Uso:
  <python-bundled> scripts/import_cbam_referencias.py --inputs-dir C:/Users/User/Downloads

Las planillas se normalizan para consulta. Los archivos originales se copian a
data/referencias/cbam y se registran con SHA-256; la aplicación no los trata
como valores calculados de una instalación.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
import sqlite3
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "sicr3p.db"
DESTINO = ROOT / "data" / "referencias" / "cbam"

FUENTES = [
    {
        "clave": "CBAM_METODOLOGIA_2025_2547",
        "archivo": "OJ_L_202502547_ES_TXT.pdf",
        "titulo": "Reglamento de Ejecución (UE) 2025/2547: métodos de cálculo de emisiones implícitas",
        "tipo": "REGLAMENTO",
        "referencia": "Reglamento de Ejecución (UE) 2025/2547",
        "version": "22-12-2025",
        "publicacion": "2025-12-22",
        "vinculante": 1,
        "url": "https://eur-lex.europa.eu/eli/reg_impl/2025/2547/oj",
        "nota": "Metodología aplicable. Las emisiones reales deben conservar evidencia por instalación y proceso.",
    },
    {
        "clave": "CBAM_VERIFICACION_2025_2546",
        "archivo": "OJ_L_202502546_ES_TXT.pdf",
        "titulo": "Reglamento de Ejecución (UE) 2025/2546: principios de verificación CBAM",
        "tipo": "REGLAMENTO",
        "referencia": "Reglamento de Ejecución (UE) 2025/2546",
        "version": "22-12-2025",
        "publicacion": "2025-12-22",
        "vinculante": 1,
        "url": "https://eur-lex.europa.eu/eli/reg_impl/2025/2546/oj",
        "nota": "Reglas de verificación, visitas y tratamiento de inexactitudes o no conformidades.",
    },
    {
        "clave": "CBAM_ACREDITACION_2025_2551",
        "archivo": "OJ_L_202502551_ES_TXT.pdf",
        "titulo": "Reglamento Delegado (UE) 2025/2551: acreditación y supervisión de verificadores",
        "tipo": "REGLAMENTO",
        "referencia": "Reglamento Delegado (UE) 2025/2551",
        "version": "22-12-2025",
        "publicacion": "2025-12-22",
        "vinculante": 1,
        "url": "https://eur-lex.europa.eu/eli/reg_del/2025/2551/oj",
        "nota": "Competencia, independencia, acreditación y supervisión del verificador.",
    },
    {
        "clave": "CBAM_GUIA_VERIFICACION_2026",
        "archivo": "GUIDANCE ON CBAM VERIFICATION AND ACCREDITATION FOR VERIFIERS AND NATIONAL ACCREDITATION BODIES (1).pdf",
        "titulo": "Guidance on CBAM verification and accreditation for verifiers and national accreditation bodies",
        "tipo": "GUIA",
        "referencia": "Comisión Europea, DG TAXUD",
        "version": "24-08-2026",
        "publicacion": "2026-08-24",
        "vinculante": 0,
        "url": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-legislation-and-guidance_en",
        "nota": "Guía explicativa no vinculante. Sus disposiciones deben contrastarse con el Diario Oficial.",
    },
    {
        "clave": "CBAM_VALORES_DEFECTO_2026_08",
        "archivo": "DV correcting act_final update_06.08.xlsx",
        "titulo": "Default values definitive period: actualización correctiva 06.08",
        "tipo": "VALOR_POR_DEFECTO",
        "referencia": "Reglamento de Ejecución (UE) 2025/2621, corregido por Reglamento de Ejecución (UE) 2026/1740",
        "version": "06-08-2026",
        "publicacion": "2026-08-10",
        "vinculante": 0,
        "url": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-legislation-and-guidance_en#default-values-and-benchmarks",
        "nota": "Planilla informativa. Los valores jurídicamente vinculantes están en los reglamentos indicados.",
    },
    {
        "clave": "CBAM_BENCHMARKS_2026_02",
        "archivo": "CBAM Benchmarks_20260206.xlsx",
        "titulo": "CBAM benchmarks definitive period",
        "tipo": "BENCHMARK",
        "referencia": "Reglamento de Ejecución (UE) 2025/2620",
        "version": "06-02-2026",
        "publicacion": "2026-02-13",
        "vinculante": 0,
        "url": "https://taxation-customs.ec.europa.eu/carbon-border-adjustment-mechanism/cbam-legislation-and-guidance_en#default-values-and-benchmarks",
        "nota": "Referencia comparativa para ajuste de asignación gratuita; no sustituye el cálculo de la instalación.",
    },
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1024 * 1024), b""):
            digest.update(bloque)
    return digest.hexdigest()


def codigo_normalizado(valor) -> str:
    return "".join(caracter for caracter in str(valor or "") if caracter.isdigit())


def numero(valor):
    if isinstance(valor, (int, float)):
        return float(valor)
    if not isinstance(valor, str):
        return None
    texto = valor.strip().replace(".", "").replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return None


def valor_texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def fila_hash(*valores) -> str:
    texto = "|".join(valor_texto(valor) for valor in valores)
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


def guardar_fuente(cursor, especificacion, origen: Path) -> int:
    DESTINO.mkdir(parents=True, exist_ok=True)
    destino = DESTINO / origen.name
    if origen.resolve() != destino.resolve():
        shutil.copy2(origen, destino)
    ruta_relativa = destino.relative_to(ROOT).as_posix()
    cursor.execute(
        """
        INSERT INTO fuentes_normativas
          (clave, titulo, tipo, jurisdiccion, referencia_legal, version, fecha_publicacion, vinculante, ruta_archivo, hash_sha256, url_oficial, nota_uso)
        VALUES (?, ?, ?, 'UE / CBAM', ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(clave) DO UPDATE SET
          titulo=excluded.titulo, tipo=excluded.tipo, referencia_legal=excluded.referencia_legal,
          version=excluded.version, fecha_publicacion=excluded.fecha_publicacion, vinculante=excluded.vinculante,
          ruta_archivo=excluded.ruta_archivo, hash_sha256=excluded.hash_sha256, url_oficial=excluded.url_oficial, nota_uso=excluded.nota_uso
        """,
        (especificacion["clave"], especificacion["titulo"], especificacion["tipo"], especificacion["referencia"],
         especificacion["version"], especificacion["publicacion"], especificacion["vinculante"], ruta_relativa,
         sha256(origen), especificacion["url"], especificacion["nota"]),
    )
    return cursor.execute("SELECT id FROM fuentes_normativas WHERE clave = ?", (especificacion["clave"],)).fetchone()[0]


def guardar_valor(cursor, fuente_id, tipo, pais, sector, codigo, descripcion, ruta, directas=None, indirectas=None, totales=None, benchmark_a=None, ruta_a=None, benchmark_b=None, ruta_b=None, bruto=""):
    codigo_limpio = valor_texto(codigo)
    hash_fila = fila_hash(fuente_id, tipo, pais, sector, codigo_limpio, descripcion, ruta, directas, indirectas, totales, benchmark_a, ruta_a, benchmark_b, ruta_b, bruto)
    cursor.execute(
        """
        INSERT INTO valores_cbam_referencia
          (fuente_id, tipo, pais, sector, codigo_cn, codigo_cn_normalizado, descripcion, ruta_produccion,
           emisiones_directas, emisiones_indirectas, emisiones_totales, benchmark_a, ruta_benchmark_a,
           benchmark_b, ruta_benchmark_b, valor_bruto, fila_hash)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(fila_hash) DO UPDATE SET
          emisiones_directas=excluded.emisiones_directas, emisiones_indirectas=excluded.emisiones_indirectas,
          emisiones_totales=excluded.emisiones_totales, benchmark_a=excluded.benchmark_a,
          benchmark_b=excluded.benchmark_b, valor_bruto=excluded.valor_bruto
        """,
        (fuente_id, tipo, pais, sector, codigo_limpio, codigo_normalizado(codigo_limpio), valor_texto(descripcion),
         valor_texto(ruta), directas, indirectas, totales, benchmark_a, valor_texto(ruta_a), benchmark_b,
         valor_texto(ruta_b), bruto, hash_fila),
    )


def importar_valores_defecto(cursor, archivo: Path, fuente_id: int) -> int:
    libro = load_workbook(archivo, read_only=True, data_only=False)
    excluidas = {"Overview", "Version History", "Annex IV"}
    total = 0
    for hoja in libro.worksheets:
        if hoja.title in excluidas:
            continue
        sector = None
        for fila in hoja.iter_rows(min_row=3, values_only=True):
            codigo = fila[0] if len(fila) > 0 else None
            descripcion = fila[1] if len(fila) > 1 else None
            if codigo is None:
                continue
            if isinstance(codigo, str) and codigo.startswith("Product CN Code"):
                continue
            if descripcion is None and all(valor is None for valor in fila[1:]):
                sector = valor_texto(codigo)
                continue
            if not codigo_normalizado(codigo):
                continue
            guardar_valor(cursor, fuente_id, "VALOR_POR_DEFECTO", hoja.title, sector, codigo, descripcion,
                          fila[5] if len(fila) > 5 else None, numero(fila[2] if len(fila) > 2 else None),
                          numero(fila[3] if len(fila) > 3 else None), numero(fila[4] if len(fila) > 4 else None),
                          bruto=" | ".join(valor_texto(valor) for valor in fila[:6]))
            total += 1
    return total


def importar_benchmarks(cursor, archivo: Path, fuente_id: int) -> int:
    hoja = load_workbook(archivo, read_only=True, data_only=False)["Benchmarks"]
    sector = None
    total = 0
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        codigo = fila[0] if len(fila) > 0 else None
        descripcion = fila[1] if len(fila) > 1 else None
        if codigo is None:
            continue
        if isinstance(codigo, str) and codigo == "CN code":
            continue
        if descripcion is None and all(valor is None for valor in fila[1:]):
            sector = valor_texto(codigo)
            continue
        if not codigo_normalizado(codigo):
            continue
        guardar_valor(cursor, fuente_id, "BENCHMARK", None, sector, codigo, descripcion, None,
                      benchmark_a=numero(fila[2] if len(fila) > 2 else None), ruta_a=fila[3] if len(fila) > 3 else None,
                      benchmark_b=numero(fila[4] if len(fila) > 4 else None), ruta_b=fila[5] if len(fila) > 5 else None,
                      bruto=" | ".join(valor_texto(valor) for valor in fila[:6]))
        total += 1
    return total


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--inputs-dir", type=Path, default=Path.home() / "Downloads")
    args = parser.parse_args()
    if not DB_PATH.exists():
        raise SystemExit(f"No existe la base de datos {DB_PATH}. Ejecuta primero npm run prebuild.")

    conexion = sqlite3.connect(DB_PATH)
    cursor = conexion.cursor()
    fuentes_ids = {}
    try:
        for fuente in FUENTES:
            origen = args.inputs_dir / fuente["archivo"]
            if not origen.exists():
                raise FileNotFoundError(f"No se encontró: {origen}")
            fuentes_ids[fuente["clave"]] = guardar_fuente(cursor, fuente, origen)
        cursor.execute("DELETE FROM valores_cbam_referencia WHERE fuente_id IN (?, ?)", (
            fuentes_ids["CBAM_VALORES_DEFECTO_2026_08"], fuentes_ids["CBAM_BENCHMARKS_2026_02"]
        ))
        valores_defecto = importar_valores_defecto(cursor, args.inputs_dir / "DV correcting act_final update_06.08.xlsx", fuentes_ids["CBAM_VALORES_DEFECTO_2026_08"])
        benchmarks = importar_benchmarks(cursor, args.inputs_dir / "CBAM Benchmarks_20260206.xlsx", fuentes_ids["CBAM_BENCHMARKS_2026_02"])
        conexion.commit()
        print(f"[cbam] {len(FUENTES)} fuentes registradas, {valores_defecto} valores por defecto y {benchmarks} benchmarks normalizados.")
    finally:
        conexion.close()


if __name__ == "__main__":
    main()
